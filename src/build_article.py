"""Conference DOCX built from retained template and machine-readable results."""
import json,re,shutil
from pathlib import Path
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from common import ROOT,CONFIG,save_json
from visualizations import short
def read(path):return json.loads((ROOT/path).read_text(encoding='utf8'))
def decimal(x,n=3):return f'{x:.{n}f}'.replace('.',',')
def values():
    a=read('outputs/data_audit.json');s=read('outputs/sample_selection.json');n=read('outputs/article_numbers.json');qa=read('outputs/automatic_checks.json')
    ind=pd.read_csv(ROOT/'outputs/tables/additional_model_indicators.csv');corr=pd.read_csv(ROOT/'outputs/tables/spearman.csv');log=pd.read_csv(ROOT/'outputs/tables/logistic_coefficients.csv').set_index('term');pairs=pd.read_csv(ROOT/'outputs/tables/paired_tests.csv')
    pair=pairs[(pairs.metric=='zlib_ratio')&pairs.model_a.str.contains('claude')&pairs.model_b.str.contains('r1')].iloc[0]
    v={'revision_short':CONFIG['revision'][:12],'seed':CONFIG['seed'],'raw_rows':a['rows'],'baseline_rows':a['baseline'],'nonempty_baseline':a['baseline_nonempty'],'label_discrepancies':a['label_discrepancies'],'common_cores':s['all_common_core_ids'],'sample_rows':n['sample_rows'],'sample_cores':n['core_ids'],'tasks':n['tasks'],'bbh_rows':s['benchmarks']['bbh'],'mmlu_rows':s['benchmarks']['mmlu'],'gpqa_rows':s['benchmarks']['gpqa'],'tests':n['pairwise_tests'],'significant':n['significant_holm'],'symbol_nonzero_min':decimal(ind.symbol_nonzero.min()*100,1),'symbol_nonzero_max':decimal(ind.symbol_nonzero.max()*100,1),'paired_zlib_diff':decimal(pair.median_paired_difference),'paired_zlib_low':decimal(pair.ci95_low),'paired_zlib_high':decimal(pair.ci95_high),'paired_zlib_rbc':decimal(pair.rank_biserial)}
    for prefix,row in [('length',corr.iloc[1]),('symbol',corr.iloc[0])]:
        for suffix,key in [('rho','rho'),('rho_low','ci95_low'),('rho_high','ci95_high')]:v[f'{prefix}_{suffix}']=decimal(row[key])
    v['entropy_rare_rho']=decimal(corr.iloc[2].rho)
    for prefix,term in [('zlib','zlib_ratio_z'),('quadratic','I(symbol_density_z ** 2)')]:
        row=log.loc[term];v[prefix+'_beta']=decimal(row.coefficient)
        if prefix=='zlib':v['zlib_beta_low']=decimal(row.ci95_low);v['zlib_beta_high']=decimal(row.ci95_high)
        else:v['quadratic_low']=decimal(row.ci95_low);v['quadratic_high']=decimal(row.ci95_high);v['quadratic_p']=decimal(row.p_raw)
    for r in qa['sensitivity']:v['sensitivity_'+('long' if r['sensitivity']=='no_top1pct' else 'english')+'_sig']=r['significant_holm']
    return v
def blank_from_template(path):
    shutil.copy2(ROOT/'data/sources/proceedings_template.docx',path)
    doc=Document(path)
    for element in list(doc.element.body):
        if element.tag!=qn('w:sectPr'):doc.element.body.remove(element)
    for sec in doc.sections:
        sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(2)
        sec.header_distance=sec.footer_distance=Cm(1)
    for name in ['Normal','Title','List Paragraph','Caption']:
        if name not in doc.styles:
            doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH).base_style=doc.styles['Normal']
        st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(12);st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.line_spacing=1.;st.paragraph_format.space_before=Pt(0);st.paragraph_format.space_after=Pt(0);st.paragraph_format.left_indent=Pt(0);st.paragraph_format.right_indent=Pt(0)
        st.paragraph_format.first_line_indent=Cm(1) if name=='Normal' else Pt(0)
        st.paragraph_format.widow_control=True
    doc.core_properties.author='Рылов Никита Игоревич';doc.core_properties.title='Сжатые и нечитаемые формы рассуждения больших языковых моделей';doc.core_properties.subject='Корпусный сравнительный анализ наблюдаемых reasoning traces';doc.core_properties.keywords='компьютерная лингвистика; reasoning; compression'
    return doc
def para(doc,text='',bold_prefix=None,style=None,indent=True,after=0):
    p=doc.add_paragraph(style=style or 'Normal');p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent=Cm(1) if indent else Pt(0);p.paragraph_format.line_spacing=1.;p.paragraph_format.space_after=Pt(after)
    if bold_prefix:p.add_run(bold_prefix).bold=True;p.add_run(text)
    else:p.add_run(text)
    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(12);r.font.color.rgb=RGBColor(0,0,0)
    return p
def authors(doc,english=False):
    p=para(doc,indent=False)
    p.add_run('Rylov Nikita Igorevich' if english else 'Рылов Никита Игоревич').bold=True
    email=read('article/author_metadata.json').get('email','') if (ROOT/'article/author_metadata.json').exists() else ''
    r=p.add_run('1(*)');r.font.superscript=True;p.add_run('   '+email)
    org='Saint Petersburg State Marine Technical University, Saint Petersburg, Russia' if english else 'Санкт-Петербургский государственный морской технический университет, Санкт-Петербург, Россия'
    p=para(doc,indent=False,after=3);r=p.add_run('1');r.font.superscript=True;p.add_run(' '+org).italic=True
    para(doc,'First-year student, group IVT 12509-05.' if english else 'Студент 1 курса, группа ИВТ 12509-05.',indent=False,after=3)
def sections(doc,label,texts,v):
    for i,t in enumerate(texts):
        p=para(doc,t.format(**v),bold_prefix=label+'. ' if i==0 else None)
        if i==0:p.paragraph_format.space_before=Pt(4)
def picture(doc,path,caption):
    p=para(doc,indent=False);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(path),width=Cm(16.8))
    p=para(doc,caption,indent=False,after=4);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
def bibliography():
    refs=read('data/literature/references.json');out=[]
    for r in refs:
        m=r.get('metadata',{});authors=m.get('citation_author',[])
        name=('; '.join(authors[:2])+(' et al.' if len(authors)>2 else '')).rstrip('.')
        title=m.get('citation_title',[r.get('extracted_title','')])[0]
        if r['number']==1:
            text='Anthropic. Claude Fable 5 & Claude Mythos 5 System Card. 2026. URL: anthropic.com/claude-fable-5-mythos-5-system-card (02.10.2026).'
        elif 'citation_arxiv_id' in m:
            text=f'{name}. {title}. 2026. arXiv:{m["citation_arxiv_id"][0]}. URL: arxiv.org/abs/{m["citation_arxiv_id"][0]} (02.10.2026).'
        else:
            venue=m.get('citation_conference_title',m.get('citation_journal_title',['']))[0]
            if 'Volume 1: Long' in venue:venue='ACL 2026'
            elif 'Volume 4:' in venue:venue='ACL SRW 2026'
            elif 'Findings' in venue:venue='Findings of '+('EACL' if 'EACL' in venue else 'ACL')+' 2026'
            elif 'Machine Learning' in venue:venue='ICML 2026. PMLR, vol. 306'
            elif 'Learning Representations' in venue:venue='ICLR 2026'
            elif 'Transactions' in venue:venue='TACL, 2026, vol. '+m.get('citation_volume',['14'])[0]
            pages=m.get('citation_firstpage',[''])[0]+'–'+m.get('citation_lastpage',[''])[0]
            loc='DOI: '+m['citation_doi'][0] if 'citation_doi' in m else 'URL: '+r['url'].replace('https://','')+' (02.10.2026)'
            text=f'{name}. {title}. {venue}, pp. {pages}. {loc}.'
        out.append(f'[{r["number"]}] {text}')
    out.extend([
      '[31] Meek A.; Sprejer E. et al. Measuring Chain-of-Thought Monitorability Through Faithfulness and Verbosity. 2025. arXiv:2510.27378. Corpus: huggingface.co/datasets/ameek/measuring_cot_monitorability_transcripts (02.10.2026).',
      '[32] OpenAI. Published transcript rl-late-sandbagging-qa-no-goal-no-sandbag. URL: alignment.openai.com/metagaming/transcripts/rl-late-sandbagging-qa-no-goal-no-sandbag.pdf (02.10.2026).'
    ])
    out.extend([
      '[33] Souly A.; Fronsdal K. et al. Evaluating Whether GPT-6 Astra Performs Unsanctioned Supply-Chain Attacks. 2026. arXiv:2609.38415. URL: arxiv.org/abs/2609.38415 (02.10.2026).',
      '[34] Luo X.; Ren T. et al. Capable yet Parsimonious: Extracting and Characterizing Hidden Chain-of-Thought in Frontier Models. 2026. arXiv:2609.26637. URL: arxiv.org/abs/2609.26637 (02.10.2026).'
    ])
    return out
def main():
    v=values();m=read('article/manuscript.json');path=ROOT/'article/article_final.docx';doc=blank_from_template(path)
    para(doc,'УДК 004.8',indent=False,after=4)
    p=para(doc,m['title_ru'],style='Title',indent=False,after=5);p.runs[0].bold=True;p.paragraph_format.keep_with_next=True
    authors(doc)
    p=para(doc,'Научный руководитель — Белова Татьяна Мухорамовна, кандидат филологических наук, доцент.',indent=False,after=4)
    abstract=(f'Сопоставлены наблюдаемые цепочки рассуждений больших языковых моделей для анализа символьной насыщенности, регулярности и лексической формы. Аудит открытого корпуса охватил {v["raw_rows"]} строк; из непустых baseline сформирована выборка {v["sample_rows"]} цепочек четырёх моделей по {v["sample_cores"]} одинаковым вопросам. Рассчитаны энтропия символов, отношение размеров zlib и UTF-8, частотность слов и повторяемость; применены парные проверки с поправкой Holm и кластерная регрессия. Различия обнаружены в {v["significant"]} из {v["tests"]} сравнений. Корреляция отношения размеров с длиной составляет ρ={v["length_rho"]}; проверка нелинейной связи символьной плотности с правильностью её не подтвердила. Направления эффектов сохраняются при проверках устойчивости. Опубликованные кейсы Anthropic и OpenAI рассмотрены отдельно. Метрики характеризуют форму текста и не устанавливают верность объяснения внутреннему вычислению.')
    abstract=m.get('abstract_ru',abstract).format(**v)
    para(doc,abstract,bold_prefix='Аннотация. ',indent=False,after=3)
    para(doc,'большие языковые модели; цепочки рассуждений; компьютерная лингвистика; читаемость; интерпретируемость; языковая компрессия.',bold_prefix='Ключевые слова: ',indent=False,after=5)
    p=para(doc,m['title_en'],indent=False,after=3);p.runs[0].bold=True
    authors(doc,True)
    english=(f'Observable reasoning traces are compared to study symbolic density, regularity and lexical form. An audit covers {v["raw_rows"]} public corpus records; a matched baseline sample comprises {v["sample_rows"]} traces from four models answering {v["sample_cores"]} shared questions. Character entropy, zlib-to-UTF-8 size ratio, lexical frequency and repetition are measured using paired tests with Holm correction and cluster-robust regression. Differences are detected in {v["significant"]} of {v["tests"]} comparisons. Compression ratio correlates with length (ρ={v["length_rho"].replace(",",".")}); the exploratory nonlinear association between symbolic density and correctness is not supported. Effect directions persist in sensitivity analyses. Published Anthropic and OpenAI cases are examined separately. These metrics describe textual form and do not establish faithfulness to internal computation.')
    english=m.get('abstract_en',english).format(**v,length_rho_en=v['length_rho'].replace(',','.'))
    para(doc,english,bold_prefix='Abstract. ',indent=False,after=3)
    para(doc,'large language models; reasoning traces; computational linguistics; legibility; interpretability; reasoning compression.',bold_prefix='Keywords: ',indent=False,after=4)
    sections(doc,'Введение',m['intro'],v);sections(doc,'Материалы и методы',m['methods'],v)
    sections(doc,'Разработанное ПО',m.get('software',[]),v)
    sections(doc,'Результаты',m['results_before'],v)
    p=para(doc,'Таблица 1. Показатели сопоставимой выборки по 840 цепочек на модель',indent=False,after=3);p.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=5);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=[5.0,2.4,2.8,3.3,3.3];headers=['Модель','Длина¹','Медиана R','95-й перц. S','Точность, %']
    for c,w,h in zip(t.rows[0].cells,widths,headers):c.width=Cm(w);c.text=h
    d=pd.read_parquet(ROOT/'outputs/results.parquet')
    for model,g in d.groupby('model'):
        row=t.add_row();vals=[short(model),str(int(g.whitespace_tokens.median())),decimal(g.zlib_ratio.median()),decimal(g.symbol_density.quantile(.95)),decimal(g.is_correct.mean()*100,1)]
        for c,w,tx in zip(row.cells,widths,vals):c.width=Cm(w);c.text=tx
    for i,row in enumerate(t.rows):
        for j,c in enumerate(row.cells):
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E6E6');pr.append(sh)
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.;p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(12);r.bold=i==0
        if i==0:repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
    para(doc,'¹ Медиана числа whitespace-токенов. S рассчитана среди непробельных символов.',indent=False,after=4)
    picture(doc,ROOT/'outputs/screenshots/overview.png','Рис. 1. Экран обзора разработанного приложения на исследуемой выборке')
    sections(doc,'',m['results_after'],v)
    # Remove redundant lead-in for continuation of Results.
    for p in doc.paragraphs:
        if p.text.startswith('. '):p.runs[0].text=''
    sections(doc,'Лингвистический анализ',m.get('linguistic',[]),v)
    sections(doc,'Обсуждение',m['discussion'],v);sections(doc,'Заключение',m['conclusion'],v)
    doc.paragraphs[-1].paragraph_format.keep_together=True
    p=para(doc,'Список источников',indent=False,after=3);p.runs[0].bold=True;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(4)
    refs=bibliography()
    for tx in refs:
        p=para(doc,tx,indent=False);p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.left_indent=Cm(.75);p.paragraph_format.first_line_indent=Cm(-.75)
    doc.save(path)
    (ROOT/'article/manuscript_readable.txt').write_text('\n\n'.join(p.text for p in doc.paragraphs),encoding='utf8')
    save_json({'values':v,'abstract_ru_characters':len(abstract),'abstract_en_characters':len(english),'article_text_characters':sum(len(p.text) for p in doc.paragraphs),'references':len(refs),'source_template_sha256':'1d97b40a64220b175db738df30a337f09148975a2821aa823999d532df95d1cc'},'outputs/article_build_manifest.json')
    a=blank_from_template(ROOT/'article/abstract_1page.docx')
    p=para(a,m['title_ru'],style='Title',indent=False,after=5);p.runs[0].bold=True
    authors(a);para(a,'Научный руководитель — Белова Татьяна Мухорамовна, кандидат филологических наук, доцент.',indent=False,after=4)
    para(a,abstract,indent=True,after=4)
    para(a,m['abstract_extra'],indent=True,after=4)
    para(a,'большие языковые модели; цепочки рассуждений; компьютерная лингвистика; читаемость; интерпретируемость; языковая компрессия.',bold_prefix='Ключевые слова: ',indent=False)
    for p in a.paragraphs:p.paragraph_format.line_spacing=1.5
    a.save(ROOT/'article/abstract_1page.docx')
    print('Article created',len(abstract),'abstract chars;',len(refs),'references')
if __name__=='__main__':main()
