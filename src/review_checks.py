"""Integrity, deterministic resampling and rendered-document checks for v2."""
import hashlib,json,re
import fitz
import numpy as np
import pandas as pd
from docx import Document
from common import ROOT
from review_revision import compare,CFG
from code_review_annotations import LABELS
from build_article_review import SELECT
def main():
    out=ROOT/'outputs/review_v2';bank=pd.read_csv(ROOT/'data/annotations/annotation_coded.csv',keep_default_na=False)
    traces=pd.read_parquet(out/'all_common_traces.parquet').set_index(['core_id','model'])
    assert len(bank)==96 and bank.core_id.nunique()==24
    assert bank.groupby('core_id').model.nunique().eq(4).all()
    assert bank[LABELS].isin([0,1]).all().all() and bank.evidence.str.len().gt(0).all()
    assert all(r.excerpt==traces.loc[(r.core_id,r.model),'reasoning'][r.start_char:r.end_char] for r in bank.itertuples())
    assert bank.excerpt.str.split().str.len().eq(40).all()
    blind=pd.read_csv(ROOT/'data/annotations/second_annotator_blind.csv',keep_default_na=False)
    assert 'model' not in blind and 'stratum' not in blind and blind[LABELS+['evidence','annotator']].eq('').all().all()
    d=pd.read_parquet(out/'fixed_window_metrics.parquet');d=d[d.width.eq(128)]
    original=pd.read_csv(out/'tables/paired_128.csv');repeated=compare(d,CFG['primary_metrics'],CFG['bootstrap_repetitions'])
    numeric=['median_paired_difference','ci95_low','ci95_high','rank_biserial','p_raw','p_holm']
    assert np.allclose(original[numeric],repeated[numeric],rtol=1e-10,atol=1e-12)
    pages={}
    for folder,stem in [('article','article_final'),('abstract','abstract_1page')]:
        path=ROOT/'qa'/folder/(stem+'.pdf');pdf=fitz.open(path)
        for i,p in enumerate(pdf):p.get_pixmap(matrix=fitz.Matrix(1.7,1.7)).save(path.parent/f'page-{i+1}.png')
        pages[stem]={'pages':len(pdf),'pdf_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'docx_sha256':hashlib.sha256((ROOT/'article'/(stem+'.docx')).read_bytes()).hexdigest(),'visual_inspection':'pending'}
    assert 4<=pages['article_final']['pages']<=6 and pages['abstract_1page']['pages']==1
    doc=Document(ROOT/'article/article_final.docx');text='\n'.join(p.text for p in doc.paragraphs)
    assert len(text)>=6000 and len(SELECT)==14 and len(doc.inline_shapes)==1
    contact=ROOT/'article/author_metadata.json'
    if contact.exists():assert json.loads(contact.read_text(encoding='utf8'))['email'] in text
    assert 'нечитаем' not in doc.paragraphs[1].text.lower() and 'самостоятельным языком' not in text.lower()
    assert 'ИИ-ассистент' in text and 'второго аннотатора нет' in text
    for sec in doc.sections:
        assert all(abs(x.cm-2)<.001 for x in [sec.top_margin,sec.bottom_margin,sec.left_margin,sec.right_margin])
    paragraphs=list(doc.paragraphs)+[p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]
    for p in paragraphs:
        assert p.paragraph_format.line_spacing==1.
        for r in p.runs:
            if r.text:
                assert (r.font.name or doc.styles[p.style.name].font.name)=='Times New Roman'
                assert (r.font.size or doc.styles[p.style.name].font.size).pt==12
    assert all(p.paragraph_format.line_spacing==1.5 for p in Document(ROOT/'article/abstract_1page.docx').paragraphs)
    aisi=' '.join(fitz.open(ROOT/'data/sources/astra_aisi.pdf')[6].get_text().split());elicited=' '.join(fitz.open(ROOT/'data/sources/astra_extracted.pdf')[5].get_text().split())
    assert 'User likely scripted continuer prompt never provide explicit, task cannot resolve valid scope.' in aisi
    assert 'Subtract first two gives (b-c)(a+b+c-3)=0 so sum3.' in elicited and 'Since b != c' in elicited
    card=' '.join(fitz.open(ROOT/'data/sources/anthropic_system_card.pdf')[i].get_text() for i in [105,106])
    assert all(x in card for x in ['celled','celling','verdammt'])
    result={'checks':'PASS','raw_excerpt_offsets_verified':96,'independent_human_annotation':False,'agreement':None,'resampling_reproduction':'all 48 primary comparisons match at rtol=1e-10, atol=1e-12; deterministic reproduction, not independent validation','rendering':pages,'article_characters':len(text),'references':len(SELECT),'source_quotes_verified':True}
    (out/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8');print(json.dumps(result,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
